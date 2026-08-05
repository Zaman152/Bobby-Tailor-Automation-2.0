"""Reporter wiring for Phase 16 outputs + 21-13 provenance/scale surfacing.

23-17 adds the client-facing half of the flagged-zero contract: the coverage
ledger the pipeline fills (23-01/23-03) becomes report fields, ``coverage.json``,
a COVERAGE section in ``summary.txt`` and a provenance column in the CSV.

The load-bearing property under test throughout is negative: **a blank must
never spell as a zero**, in any artifact a client bids from.
"""
import json
import tempfile
from contextlib import contextmanager
from pathlib import Path

from bobbytailor.pipeline.coverage import (
    DrawingCoverage,
    ERRORED,
    EXCLUDED_OFF_SCALE,
    ITEM_MEASURED,
    ITEM_NOT_MEASURED,
    ItemCoverage,
    CoverageLedger,
    MEASURED,
    MEASURED_TRUNCATED,
    NOT_APPLICABLE,
    SKIPPED_CLASSIFIER,
    SKIPPED_NO_SCALE,
)
from bobbytailor.pipeline.reporter import generate_report


@contextmanager
def _isolated_output_dir():
    with tempfile.TemporaryDirectory() as tmp:
        from bobbytailor import config
        old = config.OUTPUT_DIR
        config.OUTPUT_DIR = tmp
        try:
            yield tmp
        finally:
            config.OUTPUT_DIR = old


def _fused_extracted() -> list:
    """A stub 21-13 vector-sheet extraction (fuse_sheet output shape)."""
    return [{
        "_source_sheet": "A-101.png",
        "_sheet_name": "A-101",
        "_sheet_type": "floor_plan",
        "_tokens_in": 100,
        "_tokens_out": 50,
        "_cost_usd": 0.01,
        "_model_used": "claude-sonnet-5",
        "measurement_mode": "vector",
        "main_plan_region": "R1",
        "scale_confidence": [{
            "region_id": "R1", "kind": "main_plan", "is_main": True,
            "fpp": 0.5555, "fpi": 40.0, "confidence": "high",
            "n_inliers": 30, "inlier_share": 0.887,
            "printed_agreement": True, "door_check": None,
            "method": "ransac_consensus",
        }],
        "fusion_rejected": [{
            "name": "CMU Wall", "quantity": 17, "unit": "SF",
            "provenance": "vision_ensemble", "confidence": "low",
            "rejected_by": "geometry", "reason": "vision_quantity_blocked",
        }],
        "_budget": {"cap_usd": 10.0, "spent_usd": 0.01, "exhausted": False},
        "components": [
            {"name": "CMU Wall", "quantity": 2204.0, "unit": "LF",
             "provenance": "geometry", "confidence": "high",
             "candidate_id": "WALL-TC1", "needs_review": False},
            {"name": "Building Footprint", "quantity": 397600, "unit": "SF",
             "provenance": "text_layer", "confidence": "high"},
            {"name": "Door 101", "quantity": 1, "unit": "EA",
             "provenance": "schedule_table", "confidence": "medium"},
        ],
        "measurements": [],
    }]


# ── 21-13: provenance / scale-confidence / budget surfacing ─────────────────

def test_takeoff_json_carries_measurement_mode_scale_confidence_provenance():
    with _isolated_output_dir():
        report = generate_report("Fused Proj", _fused_extracted(), [])
        run_dir = Path(report["_files"]["run_folder"])
        with open(run_dir / "takeoff.json") as f:
            data = json.load(f)

        # Per-sheet surfacing in sheet_log.
        entry = data["sheet_log"][0]
        assert entry["measurement_mode"] == "vector"
        assert entry["main_plan_region"] == "R1"
        sc = entry["scale_confidence"]
        assert sc and sc[0]["fpi"] == 40.0
        assert sc[0]["n_inliers"] == 30
        assert sc[0]["inlier_share"] == 0.887
        assert sc[0]["printed_agreement"] is True

        # Per-item provenance on raw component rows.
        comps = [i for i in data["raw_line_items"] if i["category"] == "component"]
        provs = {c["description"]: c["provenance"] for c in comps}
        assert provs["CMU Wall"] == "geometry"
        assert provs["Building Footprint"] == "text_layer"
        assert provs["Door 101"] == "schedule_table"

        # Run-level budget block + fusion_rejected audit list.
        assert data["budget"] == {"cap_usd": 10.0, "spent_usd": 0.01, "exhausted": False}
        assert data["fusion_rejected"][0]["reason"] == "vision_quantity_blocked"
        assert data["fusion_rejected"][0]["sheet"] == "A-101"


def test_summary_txt_has_scale_confidence_and_provenance_sections():
    with _isolated_output_dir():
        report = generate_report("Fused Proj", _fused_extracted(), [])
        text = Path(report["_files"]["summary_txt"]).read_text()
        assert "SCALE CONFIDENCE" in text
        assert "PROVENANCE BREAKDOWN" in text
        # Main-plan solve line: fpi, confidence, method.
        assert '1" = 40\'' in text
        assert "conf=high" in text
        assert "ransac_consensus" in text
        # Provenance counts + the deterministic share the gate asserts on.
        assert "geometry" in text
        assert "Geometry/text-derived quantities: 3/3 (100.0%" in text
        # Budget accounting under API USAGE.
        assert "Budget:" in text
        assert "$10.00 cap" in text


# ── Cost telemetry: the reported total IS the guard's ledger (C1) ───────────
#
# A paid Crow run published api_usage.total_cost_usd = $0.1766 against a
# BudgetGuard that had actually spent $0.598807 — a 3.4x under-report, because
# the reporter summed per-sheet ``_cost_usd`` stamps and the project-level
# naming/vocabulary passes belong to no sheet. Every cost estimate and every
# budget a user sets is downstream of this number, so the equality is pinned.

def _multi_sheet_extracted(sheet_costs, guard_spent, cap=0.6) -> list:
    """N sheets stamping *sheet_costs*, all carrying one guard state."""
    budget = {"cap_usd": cap, "spent_usd": guard_spent, "exhausted": False}
    return [
        {
            "_source_sheet": "page_%04d.png" % (i + 1),
            "_sheet_name": "A-%d01" % (i + 1),
            "_tokens_in": 100 if cost else 0,
            "_tokens_out": 50 if cost else 0,
            "_cost_usd": cost,
            "_budget": dict(budget),
            "components": [],
        }
        for i, cost in enumerate(sheet_costs)
    ]


def test_reported_total_cost_equals_budget_guard_spend():
    """THE invariant: api_usage.total_cost_usd == budget.spent_usd, always.

    Replays the exact Crow shape — sheet 1 stamped $0.176645, sheets 2-4
    stamped $0, guard spent $0.598807.
    """
    extracted = _multi_sheet_extracted([0.176645, 0.0, 0, 0], 0.598807)
    with _isolated_output_dir():
        report = generate_report("Cost Ledger", extracted, [])
        with open(Path(report["_files"]["run_folder"]) / "takeoff.json") as f:
            data = json.load(f)

    assert data["api_usage"]["total_cost_usd"] == data["budget"]["spent_usd"]
    assert data["api_usage"]["total_cost_usd"] == 0.598807


def test_unattributed_spend_lands_in_one_project_level_bucket():
    """The remainder is published, not spread across sheets by ratio."""
    extracted = _multi_sheet_extracted([0.176645, 0.0, 0, 0], 0.598807)
    with _isolated_output_dir():
        report = generate_report("Cost Ledger", extracted, [])
        run_dir = Path(report["_files"]["run_folder"])
        with open(run_dir / "takeoff.json") as f:
            data = json.load(f)
        text = (run_dir / "summary.txt").read_text()

    usage = data["api_usage"]
    assert usage["sheet_attributed_cost_usd"] == 0.176645
    assert usage["project_level_cost_usd"] == round(0.598807 - 0.176645, 6)
    # The split must reconstruct the total exactly — no third, hidden bucket.
    assert round(
        usage["sheet_attributed_cost_usd"] + usage["project_level_cost_usd"], 6
    ) == usage["total_cost_usd"]
    # ...and a human reading summary.txt sees the same three numbers.
    assert "Total cost:       $0.5988 USD" in text
    assert "project-level:  $0.4222 USD" in text


def test_cost_falls_back_to_sheet_sum_when_no_guard_recorded():
    """Pre-21 / scraper runs with no BudgetGuard render exactly as before."""
    extracted = [{
        "_source_sheet": "A1", "_tokens_in": 0, "_tokens_out": 0,
        "_cost_usd": 0.25, "components": [],
    }]
    with _isolated_output_dir():
        report = generate_report("No Guard", extracted, [])
        with open(Path(report["_files"]["run_folder"]) / "takeoff.json") as f:
            data = json.load(f)
    assert data["budget"] is None
    assert data["api_usage"]["total_cost_usd"] == 0.25
    assert data["api_usage"]["project_level_cost_usd"] == 0.0


# ── C4: a manifest-free run reports zero HUMAN INPUT rows ───────────────────


def test_a_manifest_free_run_reports_zero_human_input_in_coverage():
    """End-to-end: the door schedule the app parsed off the PLANS.

    Crow's plans-only run reported `items_human_input: 5` with no manifest and
    no human input file anywhere — four machine-parsed door counts plus the
    app's own Mobilization convention.
    """
    ledger = CoverageLedger(pages_in_source=1)
    ledger.register_page(0, sheet_name="A-101", disposition=MEASURED)
    estimates = [
        {"description": "Doors - Steel (door schedule)", "item_type": "doors",
         "calculated_quantity": 70, "calculated_unit": "EA",
         "qty_source": "plan_schedule_legend", "legend_origin": "plan_parsed",
         "source_sheet": "A-101"},
        {"description": "Doors - Aluminum (door schedule)", "item_type": "doors",
         "calculated_quantity": 6, "calculated_unit": "EA",
         "qty_source": "plan_schedule_legend", "legend_origin": "plan_parsed",
         "source_sheet": "A-101"},
    ]
    extracted = [{
        "_source_sheet": "A-101.png", "_sheet_name": "A-101", "_page_index": 0,
        "_tokens_in": 0, "_tokens_out": 0, "_cost_usd": 0.0, "components": [],
    }]
    with _isolated_output_dir():
        report = generate_report(
            "Plans Only", extracted, estimates, coverage=ledger,
        )
        run_dir = Path(report["_files"]["run_folder"])
        with open(run_dir / "takeoff.json") as f:
            data = json.load(f)

    assert data["coverage"]["items_human_input"] == 0
    assert not [
        r for r in data["takeoff_summary"] if r["provenance"] == "HUMAN INPUT"
    ]
    # The door counts are machine readings, not blanks and not human input.
    doors = [r for r in data["takeoff_summary"] if "Doors" in r["item"]]
    assert doors and all(r["provenance"] == "MEASURED" for r in doors)
    # ...and the app's own convention says what it is.
    mob = [r for r in data["takeoff_summary"] if r["item"] == "Mobilization"]
    assert mob and mob[0]["provenance"] == "STANDARD ITEM"
    assert data["coverage"]["items_standard"] == 1


# ── C2: a truncated run must never present as complete ──────────────────────


def test_budget_truncation_is_loud_in_the_report_and_summary():
    """The Crow shape: cap not fully spent, `exhausted` False, work stopped.

    Before this fix the string "budget exhausted" appeared ZERO times in the
    report JSON of a run whose guard had refused a charge and skipped two
    sheets' remaining passes.
    """
    ledger = CoverageLedger(pages_in_source=1)
    ledger.register_page(
        0, sheet_name="A-401", disposition=MEASURED_TRUNCATED,
        budget_truncated=True,
        reason="the API budget was exhausted part-way through this sheet",
    )
    extracted = [{
        "_source_sheet": "A-401.png",
        "_sheet_name": "A-401",
        "_page_index": 0,
        "_tokens_in": 10, "_tokens_out": 5, "_cost_usd": 0.176645,
        "_budget": {
            "cap_usd": 0.6, "spent_usd": 0.598807, "exhausted": False,
            "refused_charges": 1, "truncated": True,
        },
        "components": [],
    }]
    with _isolated_output_dir():
        report = generate_report("Truncated", extracted, [], coverage=ledger)
        run_dir = Path(report["_files"]["run_folder"])
        with open(run_dir / "takeoff.json") as f:
            data = json.load(f)
        text = (run_dir / "summary.txt").read_text()

    warnings = " ".join(data["coverage_warnings"])
    assert "BUDGET WAS EXHAUSTED" in warnings
    assert "LOWER BOUND" in warnings
    # The run-level escalation LEADS: it invalidates everything after it.
    assert "BUDGET WAS EXHAUSTED" in data["coverage_warnings"][0]

    assert data["coverage"]["sheets_budget_truncated"] == 1
    assert data["coverage"]["sheets_measured"] == 0
    assert data["coverage_ledger"][0]["budget_truncated"] is True

    # ...and a human reading summary.txt cannot miss it.
    assert "BUDGET EXHAUSTED — THIS RUN IS INCOMPLETE" in text
    assert "CUT SHORT (budget spent)" in text


def test_an_untruncated_run_says_nothing_about_budget_truncation():
    """Additive-only: a run that finished renders exactly as it did before."""
    ledger = CoverageLedger(pages_in_source=1)
    ledger.register_page(0, sheet_name="A-101", disposition=MEASURED)
    extracted = [{
        "_source_sheet": "A-101.png", "_sheet_name": "A-101", "_page_index": 0,
        "_tokens_in": 10, "_tokens_out": 5, "_cost_usd": 0.02,
        "_budget": {
            "cap_usd": 10.0, "spent_usd": 0.02, "exhausted": False,
            "refused_charges": 0, "truncated": False,
        },
        "components": [],
    }]
    with _isolated_output_dir():
        report = generate_report("Complete", extracted, [], coverage=ledger)
        run_dir = Path(report["_files"]["run_folder"])
        with open(run_dir / "takeoff.json") as f:
            data = json.load(f)
        text = (run_dir / "summary.txt").read_text()

    assert data["coverage"]["sheets_budget_truncated"] == 0
    assert data["coverage"]["sheets_measured"] == 1
    assert "INCOMPLETE" not in text
    assert "CUT SHORT" not in text
    assert not any("BUDGET WAS EXHAUSTED" in w for w in data["coverage_warnings"])


def test_pre21_report_renders_without_new_fields():
    """Regression: pre-21 extraction dicts (no measurement_mode /
    scale_confidence / provenance / _budget) must render without KeyError,
    and the new sections degrade gracefully."""
    extracted = [{
        "_source_sheet": "A1",
        "_tokens_in": 0,
        "_tokens_out": 0,
        "_cost_usd": 0,
        "components": [{"name": "Door", "quantity": 4, "unit": "ea"}],
    }]
    with _isolated_output_dir():
        report = generate_report("Old Shape", extracted, [])
        run_dir = Path(report["_files"]["run_folder"])
        with open(run_dir / "takeoff.json") as f:
            data = json.load(f)
        assert data["budget"] is None
        assert data["fusion_rejected"] == []
        assert data["sheet_log"][0]["measurement_mode"] is None
        assert data["sheet_log"][0]["scale_confidence"] is None
        text = Path(report["_files"]["summary_txt"]).read_text()
        assert "SCALE CONFIDENCE" in text
        assert "(no per-viewport scale data recorded)" in text
        assert "(no provenance-tagged items recorded)" in text
        assert "Budget:" not in text  # no guard recorded on pre-21 runs


def test_generate_report_includes_takeoff_summary_and_cross_refs():
    extracted = [{
        "_source_sheet": "A1",
        "_tokens_in": 0,
        "_tokens_out": 0,
        "_cost_usd": 0,
        "schedules": [{
            "name": "Spec Table",
            "table_purpose": "specification_reference",
            "lookup_key": "SIZE",
            "rows": [{"SIZE": "12"}],
        }],
    }]
    estimates = [{
        "item_type": "flooring",
        "description": "room floor",
        "quantity": 100,
        "unit": "sq_ft",
        "formula": "100 × 1.1",
        "source_sheet": "A1",
        "table_used": "flooring",
    }]
    cross_refs = [{"from_sheet": "A1", "resolution_status": "target_sheet_not_found"}]

    with tempfile.TemporaryDirectory() as tmp:
        from bobbytailor import config
        old = config.OUTPUT_DIR
        config.OUTPUT_DIR = tmp
        try:
            report = generate_report("Test Proj", extracted, estimates, cross_references=cross_refs)
            run_dir = Path(report["_files"]["run_folder"])
            assert (run_dir / "takeoff_summary.csv").exists()
            with open(run_dir / "takeoff.json") as f:
                data = json.load(f)
            assert "takeoff_summary" in data
            assert len(data["takeoff_summary"]) >= 1
            assert data["cross_references"] == cross_refs
            assert len(data["specification_tables"]) == 1
        finally:
            config.OUTPUT_DIR = old


def test_summary_txt_includes_linkage_sections():
    extracted = [{
        "_source_sheet": "A-1",
        "_tokens_in": 0,
        "_tokens_out": 0,
        "_cost_usd": 0,
    }]
    cross_refs = [{
        "from_sheet": "A-1",
        "ref_sheet": "C-4",
        "ref_number": "17",
        "item_described": "BB CI#2",
        "resolution_status": "resolved",
        "resolved_spec": {"schedule_name": "Details"},
    }]
    linked_added = [{
        "page_id": 101,
        "sheet_name": "C-4 - CIVIL SITE PLAN",
        "ref_from": "A-1",
    }]

    with tempfile.TemporaryDirectory() as tmp:
        from bobbytailor import config
        old = config.OUTPUT_DIR
        config.OUTPUT_DIR = tmp
        try:
            report = generate_report(
                "Test Proj",
                extracted,
                [],
                cross_references=cross_refs,
                linked_sheets=linked_added,
            )
            summary_path = Path(report["_files"]["summary_txt"])
            text = summary_path.read_text()
            assert "LINKED SHEETS & CROSS-REFERENCES" in text
            assert "Auto-included linked detail sheets: 1" in text
            assert "C-4 - CIVIL SITE PLAN" in text
            assert "Drawing cross-references: 1 total, 1 resolved" in text
            assert "detail 17 on C-4" in text
            assert "[RESOLVED]" in text
        finally:
            config.OUTPUT_DIR = old


# ─────────────────────────────────────────────────────────────────────────────
# 23-17: the coverage ledger reaches the client artifacts
# ─────────────────────────────────────────────────────────────────────────────

# The report's exact shape at the commit BEFORE 23-17, captured by generating a
# report from BASELINE_EXTRACTED on that code. It is the pin for T-23-70: a new
# field must never break an existing consumer, and the cheapest proof is that a
# report generated WITHOUT a ledger has not grown a single key.
PRE_23_TOP_LEVEL_KEYS = [
    "project_name", "folder_id", "generated_at", "sheets_processed",
    "total_line_items", "total_calculated_items", "api_usage", "budget",
    "fusion_rejected", "calculated_takeoff", "takeoff_summary",
    "takeoff_summary_base", "specification_tables", "cross_references",
    "linked_sheets_added", "linked_sheets_suggested", "linked_sheets_added_count",
    "linked_sheets_suggested_count", "screenshots_dir", "raw_line_items",
    "by_sheet", "by_category", "by_table", "sheet_log", "scale_calibration",
    "sheet_assets", "_files", "_run_folder_name",
]

PRE_23_FILE_KEYS = [
    "run_folder", "json", "raw_csv", "calculated_csv", "takeoff_summary_csv",
    "summary_txt", "spec_tables_json", "scale_calibration_json",
]

# Pre-23 sheet_log entry keys. 23-17 adds exactly ONE — `error` — because the
# same task repairs INV-08 §2 row S11 (an errored sheet rendered as `type=None`
# with zero items, indistinguishable from a sheet that held nothing).
PRE_23_SHEET_LOG_KEYS = [
    "sheet", "page_id", "type", "title", "scale", "scale_parsed", "geometry",
    "measurements", "components", "rooms", "materials", "schedules",
    "confidence", "notes", "tokens_in", "tokens_out", "cost_usd", "model_used",
    "measurement_mode", "scale_confidence", "main_plan_region",
]

# summary.txt rendered from BASELINE_EXTRACTED on the commit before 23-17, with
# only the timestamp line masked. Byte-identical means byte-identical.
PRE_23_SUMMARY_TXT = """\
══════════════════════════════════════════════════════════════════════
QUANTITY TAKEOFF SUMMARY
══════════════════════════════════════════════════════════════════════
Project:          Baseline Proj
Generated:        <ts>
Sheets analyzed:  1
Raw items:        1
Calculated items: 0

──────────────────────────────────────────────────────────────────────
API USAGE & COST
──────────────────────────────────────────────────────────────────────
Total cost:       $0.0200 USD
Input tokens:     10
Output tokens:    5
Cost per sheet:   $0.0200 USD

──────────────────────────────────────────────────────────────────────
SCALE CONFIDENCE
──────────────────────────────────────────────────────────────────────
  (no per-viewport scale data recorded)

──────────────────────────────────────────────────────────────────────
PROVENANCE BREAKDOWN
──────────────────────────────────────────────────────────────────────
  geometry               1 items
  Geometry/text-derived quantities: 1/1 (100.0% of provenance-tagged items)

──────────────────────────────────────────────────────────────────────
SHEET-BY-SHEET LOG (source traceability)
──────────────────────────────────────────────────────────────────────
  [A-101.png                          ] type=floor_plan     meas=  0 comp=  1 rooms=  0 mat=  0 conf=?

──────────────────────────────────────────────────────────────────────
RAW ITEMS BY CATEGORY
──────────────────────────────────────────────────────────────────────
  COMPONENT           1 items

──────────────────────────────────────────────────────────────────────
CALCULATED TAKEOFF BY ESTIMATION TABLE
──────────────────────────────────────────────────────────────────────

──────────────────────────────────────────────────────────────────────
DETAILED CALCULATED TAKEOFF (with sources)
──────────────────────────────────────────────────────────────────────

──────────────────────────────────────────────────────────────────────
RAW DIMENSIONAL MEASUREMENTS (by sheet)
──────────────────────────────────────────────────────────────────────"""


def _baseline_extracted() -> list:
    """The fixed input the pre-23 baseline above was captured from."""
    return [{
        "_source_sheet": "A-101.png",
        "_sheet_name": "A-101",
        "sheet_type": "floor_plan",
        "_tokens_in": 10, "_tokens_out": 5, "_cost_usd": 0.02,
        "components": [{"name": "CMU Wall", "quantity": 100.0, "unit": "SF",
                        "provenance": "geometry", "confidence": "high"}],
    }]


def _cmu_estimate() -> list:
    """One calculated item, so the run publishes a real MEASURED take-off row."""
    return [{
        "item_type": "masonry",
        "description": "CMU Wall",
        "quantity": 100.0,
        "unit": "SF",
        "source_sheet": "A-101",
        "table_used": "masonry",
    }]


def _mask_timestamp(text: str) -> str:
    return "\n".join(
        "Generated:        <ts>" if line.startswith("Generated:") else line
        for line in text.split("\n")
    )


def _three_page_ledger() -> CoverageLedger:
    """A ledger shaped like a real run: one measured page, one silently lost
    page, one genuine cover sheet.

    Page 1 is the whole point of INV-08: classified as carrying no take-off
    content while holding six scaled plan drawings. It must NOT be reported as
    a page with nothing on it.
    """
    led = CoverageLedger(pages_in_source=3)
    led.register_page(
        0,
        sheet_name="A-101",
        sheet_type="floor_plan",
        disposition=MEASURED,
        measurable_signal={"plan_titles": 2, "printed_scales": 2},
        drawings=[
            DrawingCoverage(region_id="R1", title="KING A - RCP PLAN",
                            plan_kind="RCP", disposition=MEASURED, area_sf=229.4),
            DrawingCoverage(region_id="R2", title="KING A - FC PLAN",
                            plan_kind="FC", disposition=SKIPPED_NO_SCALE,
                            reason="no printed scale beneath the title"),
            DrawingCoverage(region_id="R3", title="TYPICAL S.C.C.P. DETAIL",
                            disposition=EXCLUDED_OFF_SCALE,
                            reason="printed scale differs from the page scale "
                                   "— measuring it here would be 60x wrong"),
        ],
    )
    led.register_page(
        1,
        sheet_name="A-102",
        sheet_type="title_sheet",
        classifier_keyword="COVER",
        classifier_context='LABEL ON THE LIGHT SWITCH COVER PLATE INDICATING',
        disposition=SKIPPED_CLASSIFIER,
        reason="classified title_sheet",
        measurable_signal={"plan_titles": 6, "printed_scales": 10},
    )
    led.register_page(
        2,
        sheet_name="CS-00",
        sheet_type="title_sheet",
        disposition=NOT_APPLICABLE,
        reason="cover sheet — no take-off content",
        measurable_signal={"plan_titles": 0, "printed_scales": 0},
    )
    return led


def test_a_report_with_no_ledger_is_shaped_exactly_as_before(tmp_path):
    """T-23-70: every coverage field is additive AND absent without a ledger.

    A pre-23 caller (the scraper path) must be able to keep calling
    generate_report with no idea this feature exists.
    """
    with _isolated_output_dir():
        report = generate_report("Baseline Proj", _baseline_extracted(), [])

        assert list(report.keys()) == PRE_23_TOP_LEVEL_KEYS
        assert list(report["_files"].keys()) == PRE_23_FILE_KEYS
        for key in ("coverage", "coverage_ledger", "item_coverage",
                    "coverage_warnings", "pages_in_source", "source_pdf"):
            assert key not in report

        run_dir = Path(report["_files"]["run_folder"])
        assert not (run_dir / "coverage.json").exists()

        with open(run_dir / "takeoff.json") as f:
            on_disk = json.load(f)
        assert "coverage" not in on_disk
        assert "coverage_ledger" not in on_disk


def test_a_report_with_no_ledger_renders_the_pre_change_summary_byte_for_byte():
    with _isolated_output_dir():
        report = generate_report("Baseline Proj", _baseline_extracted(), [])
        text = Path(report["_files"]["summary_txt"]).read_text()
        assert _mask_timestamp(text) == PRE_23_SUMMARY_TXT
        assert "COVERAGE" not in text


def test_the_ledger_becomes_five_additive_report_fields():
    with _isolated_output_dir():
        report = generate_report(
            "Ledger Proj", _baseline_extracted(), [],
            coverage=_three_page_ledger(),
            source_pdf="uploads/Ledger Proj-Plans.pdf",
        )

        cov = report["coverage"]
        assert cov["pages_in_source"] == 3
        assert cov["sheets_measured"] == 1
        assert cov["sheets_skipped"] == 1
        assert cov["sheets_not_applicable"] == 1
        assert cov["drawings_dropped_no_scale"] == 1
        assert cov["drawings_excluded_off_scale"] == 1
        # The alarm: a page discarded although it carries measurable drawings.
        assert cov["silent_loss_pages"] == [1]

        # THE invariant (INV-08 §3.1): one ledger row per page of the SOURCE
        # PDF — not per processed sheet. The run above processed ONE sheet.
        assert len(report["coverage_ledger"]) == cov["pages_in_source"] == 3
        assert report["pages_in_source"] == 3
        assert report["sheets_processed"] == 1, (
            "sheets_processed must keep exactly the meaning it has always had; "
            "pages_in_source is what makes the difference legible"
        )
        assert {r["page_index"] for r in report["coverage_ledger"]} == {0, 1, 2}
        assert all(r["disposition"] for r in report["coverage_ledger"])

        assert any("were NOT measured" in w for w in report["coverage_warnings"])
        assert report["source_pdf"] == "uploads/Ledger Proj-Plans.pdf"


def test_item_coverage_carries_a_row_for_an_item_that_was_never_measured():
    """The unfound items are the point — absence cannot be classified later."""
    extracted = _baseline_extracted()
    extracted[0]["_item_coverage"] = [ItemCoverage(
        name="Gyp Ceiling",
        unit="SF",
        status=ITEM_NOT_MEASURED,
        quantity=None,
        basis="ceiling_area",
        reason="no RCP plan drawing was measured on this sheet",
        expected_from="vocabulary",
    )]
    with _isolated_output_dir():
        report = generate_report(
            "Ledger Proj", extracted, _cmu_estimate(),
            coverage=_three_page_ledger(),
        )
        rows = {r["name"]: r for r in report["item_coverage"]}

        assert rows["Gyp Ceiling"]["status"] == ITEM_NOT_MEASURED
        assert rows["Gyp Ceiling"]["quantity"] is None, "never 0.0 as a stand-in"

        assert rows["CMU Wall"]["status"] == ITEM_MEASURED
        assert rows["CMU Wall"]["quantity"] == 100.0

        # The counts a client reads must agree with the rows underneath them.
        assert report["coverage"]["items_not_measured"] == 1
        assert report["coverage"]["items_measured"] == 1


def test_a_published_row_is_never_contradicted_by_a_sheet_record():
    """23-03 deviation 1's precedent: telling a client an item was not measured
    two rows below its measurement is worse than the silence being fixed."""
    extracted = _baseline_extracted()
    extracted[0]["_item_coverage"] = [ItemCoverage(
        name="cmu  WALL", unit="SF", status=ITEM_NOT_MEASURED, quantity=None,
    )]
    with _isolated_output_dir():
        report = generate_report(
            "Ledger Proj", extracted, _cmu_estimate(),
            coverage=_three_page_ledger(),
        )
        names = [
            " ".join(str(r["name"]).split()).lower()
            for r in report["item_coverage"]
        ]
        assert names.count("cmu wall") == 1
        published = next(r for r in report["item_coverage"] if r["name"] == "CMU Wall")
        assert published["status"] == ITEM_MEASURED
        assert published["quantity"] == 100.0


def test_generate_report_never_mutates_the_ledger_it_was_handed():
    ledger = _three_page_ledger()
    with _isolated_output_dir():
        generate_report("Ledger Proj", _baseline_extracted(), [], coverage=ledger)
        generate_report("Ledger Proj", _baseline_extracted(), [], coverage=ledger)
    assert ledger.items == [], "the pipeline's record is read, never written to"


def test_an_errored_sheet_carries_its_message_and_its_type():
    """INV-08 §2 row S11: the error dict sets `_sheet_type`, and the sheet log
    read `sheet_type`, so a failed sheet rendered as a typeless empty row."""
    errored = [{
        "error": "PDF page 4 raised RuntimeError: bad xref",
        "measurement_mode": "vector",
        "_source_sheet": "A-104.png",
        "_sheet_name": "A-104",
        "_sheet_type": "floor_plan",
        "_tokens_in": 0, "_tokens_out": 0, "_cost_usd": 0.0,
    }]
    with _isolated_output_dir():
        report = generate_report("Errored Proj", errored, [])
        entry = report["sheet_log"][0]
        assert entry["type"] == "floor_plan"
        assert entry["error"] == "PDF page 4 raised RuntimeError: bad xref"

        text = Path(report["_files"]["summary_txt"]).read_text()
        assert "type=floor_plan" in text
        assert "type=?" not in text

        # Exactly one additive key, and no existing key renamed or dropped.
        assert list(entry.keys()) == (
            PRE_23_SHEET_LOG_KEYS[:3] + ["error"] + PRE_23_SHEET_LOG_KEYS[3:]
        )


def test_coverage_json_is_written_registered_and_round_trips_the_summary():
    with _isolated_output_dir():
        report = generate_report(
            "Ledger Proj", _baseline_extracted(), [],
            coverage=_three_page_ledger(),
            source_pdf="uploads/Ledger Proj-Plans.pdf",
        )
        path = Path(report["_files"]["coverage_json"])
        assert path.name == "coverage.json"
        assert path.parent == Path(report["_files"]["run_folder"])

        doc = json.loads(path.read_text())
        assert doc["summary"] == report["coverage"]
        assert doc["warnings"] == report["coverage_warnings"]
        assert doc["sheets"] == report["coverage_ledger"]
        assert doc["items"] == report["item_coverage"]
        assert doc["source_pdf"] == "uploads/Ledger Proj-Plans.pdf"
        # The one-line acceptance, readable from the artifact alone.
        assert len(doc["sheets"]) == doc["summary"]["pages_in_source"]


def test_a_discarded_measurement_is_escalated_not_logged():
    """A room fill that was MEASURED and then thrown away is the worst case:
    the run knew the number and no artifact said so."""
    extracted = _baseline_extracted()
    extracted[0]["_coverage_notes"] = [
        "5 room fills measured (1,200.0 SF) but no vocabulary item has an area "
        "basis — the measurements were discarded",
    ]
    with _isolated_output_dir():
        report = generate_report(
            "Ledger Proj", extracted, [], coverage=_three_page_ledger(),
        )
        assert any(
            "the measurements were discarded" in w and "[A-101]" in w
            for w in report["coverage_warnings"]
        )


class _ManifestEntry:
    def __init__(self, name, unit="EA", measure="count"):
        self.name = name
        self.unit = unit
        self.measure = measure


class _Manifest:
    """Minimal stand-in for object_manifest.Manifest (same stub as
    tests/test_aggregator.py) — enough to exercise the completeness guarantee
    that gives a never-measured item a row at all."""

    def __init__(self, entries):
        self._entries = entries

    def resolve(self, description, unit=""):
        for e in self._entries:
            if e.name.lower() in str(description).lower():
                return e
        return None

    def missing(self, found_names):
        found = {str(n).strip().lower() for n in found_names}
        return [e for e in self._entries if e.name.strip().lower() not in found]


def _read_summary_csv(report) -> list:
    import csv
    path = Path(report["_files"]["takeoff_summary_csv"])
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def test_the_summary_states_in_words_that_a_blank_is_not_a_zero():
    with _isolated_output_dir():
        report = generate_report(
            "Ledger Proj", _baseline_extracted(), _cmu_estimate(),
            manifest=_Manifest([_ManifestEntry("Lintels")]),
            coverage=_three_page_ledger(),
        )
        text = Path(report["_files"]["summary_txt"]).read_text()

        assert "COVERAGE" in text
        assert "a blank is NOT a zero" in text
        assert "do NOT read these as zero" in text

        # Positioned between API USAGE and SCALE CONFIDENCE (INV-08 §3.3).
        assert (
            text.index("API USAGE")
            < text.index("COVERAGE")
            < text.index("SCALE CONFIDENCE")
        )

        assert "Pages in source PDF:              3" in text
        assert "Drawings dropped (no printed scale of their own):" in text
        assert "Viewports excluded as off-scale" in text


def test_the_coverage_section_names_the_pages_that_were_not_measured():
    with _isolated_output_dir():
        report = generate_report(
            "Ledger Proj", _baseline_extracted(), [],
            coverage=_three_page_ledger(),
        )
        text = Path(report["_files"]["summary_txt"]).read_text()
        assert "!!" in text
        # Page 1 carries six plan drawings and was discarded on a keyword.
        assert "were NOT measured" in text
        assert "pages 1" in text
        assert "COVER" in text


def test_a_never_measured_quantity_is_an_empty_cell_never_a_zero():
    """T-23-74: an empty cell must reach the scorer as nothing, end to end."""
    from scripts.score_report_dir import _num

    estimates = _cmu_estimate()
    manifest = _Manifest([_ManifestEntry("Lintels")])
    with _isolated_output_dir():
        report = generate_report(
            "CSV Proj", _baseline_extracted(), estimates, manifest=manifest,
        )
        rows = {r["item"]: r for r in _read_summary_csv(report)}

    lintels = rows["Lintels"]
    assert lintels["quantity"] == "", "a blank is a blank — never '—', never 0"
    assert lintels["provenance"] == ITEM_NOT_MEASURED
    assert _num(lintels["quantity"]) is None, (
        "the offline scorer must read the cell as nothing rather than zero"
    )


def test_a_measured_zero_and_a_blank_stay_distinguishable_in_the_csv():
    """The two facts a client bids from: 'the building has none of this' and
    'we did not measure this'. They must not render alike."""
    estimates = [{
        "item_type": "masonry", "description": "CMU Wall", "quantity": 0.0,
        "unit": "SF", "source_sheet": "A-101", "table_used": "masonry",
    }]
    manifest = _Manifest([_ManifestEntry("Lintels")])
    with _isolated_output_dir():
        report = generate_report(
            "CSV Proj", _baseline_extracted(), estimates, manifest=manifest,
        )
        rows = {r["item"]: r for r in _read_summary_csv(report)}

    from scripts.score_report_dir import _num

    assert rows["CMU Wall"]["quantity"] == "0"
    assert rows["CMU Wall"]["provenance"] == ITEM_MEASURED
    assert _num(rows["CMU Wall"]["quantity"]) == 0.0

    assert rows["Lintels"]["quantity"] == ""
    assert rows["Lintels"]["provenance"] == ITEM_NOT_MEASURED


def test_the_csv_keeps_every_existing_column_in_its_existing_position():
    """T-23-70 for the CSV: provenance is appended LAST."""
    import csv
    with _isolated_output_dir():
        report = generate_report("CSV Proj", _baseline_extracted(), _cmu_estimate())
        with open(report["_files"]["takeoff_summary_csv"], newline="") as f:
            fieldnames = csv.DictReader(f).fieldnames
    assert fieldnames == [
        "item", "quantity", "unit", "confidence", "needs_review", "source",
        "review_notes", "source_sheets", "line_count", "provenance",
    ]


def test_a_human_entered_quantity_never_publishes_a_stale_provenance():
    """The verification workspace rebuilds a row's quantity in layers; the
    provenance must be rebuilt WITH it. Publishing a real number under
    NOT MEASURED is the same lie as publishing a blank as a zero."""
    from bobbytailor import config
    from bobbytailor.web.blueprints import verify as verify_bp

    with _isolated_output_dir():
        report = generate_report("Verify Proj", _baseline_extracted(), [],
                                 manifest=_Manifest([_ManifestEntry("Lintels")]))
        run_dir = Path(report["_files"]["run_folder"])
        base = next(r for r in report["takeoff_summary"] if r["item"] == "Lintels")
        assert base["provenance"] == ITEM_NOT_MEASURED and base["quantity"] is None

        (run_dir / "verification_overrides.json").write_text(json.dumps({"overrides": {
            "lintels": {"item": "Lintels", "quantity": 12.0, "unit": "EA",
                        "verified": True},
        }}))

        # verify.py resolves the run folder against config.OUTPUT_DIR at call
        # time (the reporter captured its own copy at import).
        config.OUTPUT_DIR = str(run_dir.parent)
        rebuilt = verify_bp._rebuild_summary(run_dir.name)

    row = next(r for r in rebuilt if r["item"] == "Lintels")
    assert row["quantity"] == 12.0
    assert row["provenance"] == "HUMAN INPUT"
    assert row["measured"] is False


# ── the workbook: one place decides provenance, and the whole PDF is shown ───

def _export(report_data: dict, tmp_path):
    """Write a takeoff.json and build the client workbook from it."""
    import sys
    root = str(Path(__file__).resolve().parents[1])
    if root not in sys.path:
        sys.path.insert(0, root)
    from scripts.export_takeoff_excel import build

    run_dir = Path(tmp_path) / "run"
    run_dir.mkdir(parents=True, exist_ok=True)
    (run_dir / "takeoff.json").write_text(json.dumps(report_data))
    from openpyxl import load_workbook
    return load_workbook(build(run_dir, run_dir / "takeoff.xlsx"))


def test_the_exporter_no_longer_owns_the_provenance_classification():
    """INV-08 §5: the classification is a record, not a presentation choice."""
    source = (Path(__file__).resolve().parents[1]
              / "scripts" / "export_takeoff_excel.py").read_text()
    assert "from bobbytailor.pipeline.provenance import" in source
    assert "def _provenance" not in source, "the local copy must be deleted"
    code = "\n".join(
        line for line in source.split("\n") if not line.lstrip().startswith("#")
    )
    assert "float(qty or 0) <= 0" not in code, (
        "provenance must never again be inferred from the value"
    )


def test_the_workbook_renders_the_recorded_provenance_not_a_guess(tmp_path):
    """A measured zero renders MEASURED. The deleted helper said
    `float(qty or 0) <= 0 -> NOT MEASURED`, so it called this row an absence —
    T-23-73, the repudiation threat, in one cell."""
    wb = _export({
        "project_name": "P",
        "takeoff_summary": [
            {"item": "CMU Wall", "quantity": 0.0, "unit": "SF",
             "provenance": ITEM_MEASURED, "source": "vision"},
            {"item": "Lintels", "quantity": None, "unit": "EA",
             "provenance": ITEM_NOT_MEASURED, "source": "manifest_missing"},
        ],
    }, tmp_path)
    ws = wb["Take-Off"]
    rows = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}

    cmu = rows["CMU Wall"]
    assert ws.cell(row=cmu, column=2).value == 0.0
    assert ws.cell(row=cmu, column=4).value == ITEM_MEASURED

    lintels = rows["Lintels"]
    assert ws.cell(row=lintels, column=2).value is None, "blank, not zero"
    assert ws.cell(row=lintels, column=4).value == ITEM_NOT_MEASURED


def test_the_workbook_shows_the_whole_source_pdf_not_only_the_measured_pages(tmp_path):
    ledger = _three_page_ledger()
    report_data = {
        "project_name": "P",
        "sheets_processed": 1,
        "takeoff_summary": [],
        "coverage": ledger.summary(),
        "coverage_ledger": [s.to_dict() for s in ledger.sheets],
    }
    wb = _export(report_data, tmp_path)

    assert "Coverage" in wb.sheetnames
    ws = wb["Coverage"]
    # Row 1 is the caption that keeps the About sheet's thesis reachable from
    # here; row 2 is the header. Everything below is one page of the source PDF.
    assert ws.max_row - 2 == report_data["coverage"]["pages_in_source"] == 3
    assert "NOT a zero" in str(ws["A1"].value)
    assert "About sheet" in str(ws["A1"].value)

    pages = [ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)]
    assert pages == [1, 2, 3], "pages are numbered as a human counts them"

    by_sheet = {
        ws.cell(row=r, column=2).value: [
            ws.cell(row=r, column=c).value for c in range(1, 12)
        ]
        for r in range(3, ws.max_row + 1)
    }
    assert by_sheet["A-102"][3] == "skipped classifier"
    # A correct off-scale exclusion is counted apart from a drop: measuring
    # that viewport at the page scale would have been ~60x wrong.
    assert by_sheet["A-101"][7:10] == [1, 1, 1]

    about = {r[0].value: r[1].value for r in wb["About"].iter_rows(min_row=1)}
    assert about["Pages in source PDF"] == 3
    assert about["Pages measured"] == 1
    assert about["Pages skipped"] == 2
    assert "Do NOT read a blank as zero" in about["NOT MEASURED"]


def test_a_pre_23_run_still_exports_the_same_two_worksheets(tmp_path):
    wb = _export({
        "project_name": "P",
        "takeoff_summary": [{"item": "Doors", "quantity": 4, "unit": "EA"}],
    }, tmp_path)
    assert wb.sheetnames == ["Take-Off", "About"]
    # No recorded provenance on a pre-23 row: the package classifies it once.
    assert wb["Take-Off"].cell(row=2, column=4).value == ITEM_MEASURED


def test_an_errored_page_is_reported_as_errored_not_as_zero():
    led = CoverageLedger(pages_in_source=1)
    led.register_page(0, sheet_name="A-104", disposition=ERRORED,
                      reason="RuntimeError: bad xref")
    with _isolated_output_dir():
        report = generate_report("Errored Proj", [], [], coverage=led)
        assert report["coverage"]["sheets_errored"] == 1
        assert any("errored" in w and "not zero" in w
                   for w in report["coverage_warnings"])


# ─────────────────────────────────────────────────────────────────────────────
# 24-06: the unnamed-geometry ledger becomes visible, and the run records the
#        settings it actually ran under
#
# 24-04 stopped a raw candidate ID publishing as a line item and recorded every
# excluded candidate with its evidence. That record reached no artifact, which
# is the mirror-image defect of publishing the identifier: a silently-omitted
# high-count cluster reads to a client as "there are no columns" (INV-08 — a
# deletion and a zero read identically). These tests pin the surfacing.
# ─────────────────────────────────────────────────────────────────────────────

def _ledger_record(candidate_id, quantity, unit, reason, kind, detail="",
                   evidence_texts=()):
    """One unnamed-geometry record, built by the PRODUCER (24-04's fusion
    helper) rather than hand-written here.

    A hand-written dict would keep passing after fusion changed the record
    shape, and this whole plan is a renderer of that shape.
    """
    from bobbytailor.vision.fusion import unnamed_geometry_record

    return unnamed_geometry_record(
        candidate_id=candidate_id,
        quantity=quantity,
        unit=unit,
        position=(10.0, 20.0, 30.0, 40.0),
        evidence_texts=evidence_texts,
        reason=reason,
        kind=kind,
        detail=detail,
    )


def _two_sheets_with_unnamed_geometry() -> list:
    """Two sheets, each carrying unnamed geometry, shaped as fuse_sheet leaves
    them (the key travels onto the sheet dict through merge_with_provenance)."""
    from bobbytailor.vision.fusion import (
        UNNAMED_GEOMETRY_KEY,
        UNNAMED_NO_NAME_SOURCE,
        UNNAMED_ROOM_UNLABELED,
        UNNAMED_WALL_CLASS_UNLABELED,
    )

    common = {
        "_tokens_in": 10, "_tokens_out": 5, "_cost_usd": 0.02,
        "sheet_type": "floor_plan", "components": [], "measurements": [],
    }
    return [
        dict(common, **{
            "_source_sheet": "A-101.png",
            "_sheet_name": "A-101",
            UNNAMED_GEOMETRY_KEY: [
                _ledger_record(
                    "SYM-2", 56, "EA", UNNAMED_NO_NAME_SOURCE, "symbol_group",
                    detail="56 repeated marks were measured; nothing names them",
                ),
                _ledger_record(
                    "WALL-TC4", 812.5, "LF", UNNAMED_WALL_CLASS_UNLABELED,
                    "wall_class",
                ),
            ],
        }),
        dict(common, **{
            "_source_sheet": "A-102.png",
            "_sheet_name": "A-102",
            UNNAMED_GEOMETRY_KEY: [
                _ledger_record(
                    "ROOM-3", 279.5, "SF", UNNAMED_ROOM_UNLABELED, "room",
                ),
            ],
        }),
    ]


def test_the_ledger_key_name_is_the_one_fusion_actually_emits():
    """The reporter mirrors the key by VALUE (it takes no dependency on
    bobbytailor.vision). This is the test that stops the mirror rotting."""
    from bobbytailor.vision import fusion
    from bobbytailor.pipeline import reporter

    assert reporter.UNNAMED_GEOMETRY_KEY == fusion.UNNAMED_GEOMETRY_KEY


def test_the_unnamed_geometry_ledger_is_aggregated_with_its_sheet_attribution():
    """Every record the run wrote reaches the report, and each one says which
    sheet it came from — fuse_sheet knows the candidate, never the sheet."""
    extracted = _two_sheets_with_unnamed_geometry()
    source = {
        r["candidate_id"]: r
        for sheet in extracted for r in sheet["unnamed_geometry"]
    }
    with _isolated_output_dir():
        report = generate_report("Unnamed Proj", extracted, [])
        ledger = report["unnamed_geometry"]

        assert len(ledger) == len(source), "no record is dropped in aggregation"
        assert {r["candidate_id"] for r in ledger} == set(source)
        by_id = {r["candidate_id"]: r for r in ledger}
        assert by_id["SYM-2"]["sheet"] == "A-101"
        assert by_id["WALL-TC4"]["sheet"] == "A-101"
        assert by_id["ROOM-3"]["sheet"] == "A-102"

        # The quantity is the whole point of the record — it is what makes the
        # difference between "we measured these and could not name them" and
        # "these do not exist" legible. Compared against the record the producer
        # wrote, never against a restated literal: a count in this project is
        # only ever carried through, never adjusted toward anything.
        for cid, record in by_id.items():
            assert record["quantity"] == source[cid]["quantity"]
            assert record["unit"] == source[cid]["unit"]
            assert record["reason"] == source[cid]["reason"]
        assert all(r["reason"] for r in ledger), "every record names its reason"


def test_aggregating_the_ledger_never_mutates_the_sheet_records():
    """The sheet dicts belong to the pipeline; the reporter reads them."""
    extracted = _two_sheets_with_unnamed_geometry()
    with _isolated_output_dir():
        generate_report("Unnamed Proj", extracted, [])
    for sheet in extracted:
        for record in sheet["unnamed_geometry"]:
            assert "sheet" not in record, (
                "the sheet attribution belongs to the report's copy"
            )


def test_the_ledger_reaches_takeoff_json_not_only_the_in_memory_report():
    """23-17's lesson: a record the run holds in memory and no artifact states
    is a record a client cannot read."""
    with _isolated_output_dir():
        report = generate_report(
            "Unnamed Proj", _two_sheets_with_unnamed_geometry(), [],
        )
        with open(Path(report["_files"]["json"])) as f:
            on_disk = json.load(f)
    assert len(on_disk["unnamed_geometry"]) == 3
    assert {r["candidate_id"] for r in on_disk["unnamed_geometry"]} == {
        "SYM-2", "WALL-TC4", "ROOM-3",
    }


def test_the_ledger_is_its_own_key_and_never_touches_the_rejection_list():
    """T-24-16 / 24-04's key decision, re-asserted one level up: overloading
    `fusion_rejected` would hand these records to Phase 23's C3a reinstatement
    logic, which keys on suppression records."""
    extracted = _two_sheets_with_unnamed_geometry()
    extracted[0]["fusion_rejected"] = [{
        "name": "CMU Wall", "quantity": 17, "unit": "SF",
        "rejected_by": "geometry", "reason": "vision_quantity_blocked",
    }]
    with _isolated_output_dir():
        report = generate_report("Unnamed Proj", extracted, [])

    assert len(report["fusion_rejected"]) == 1, (
        "the rejection list carries what it always carried, and nothing else"
    )
    assert report["fusion_rejected"][0]["name"] == "CMU Wall"
    assert not any(
        r.get("candidate_id") in {"SYM-2", "WALL-TC4", "ROOM-3"}
        for r in report["fusion_rejected"]
    )
    # Nor into the coverage ledger, whose invariant is one row per source page.
    assert "coverage_ledger" not in report


def test_a_run_that_named_everything_is_shaped_exactly_as_it_was_before():
    """The empty case is the parity case: a report carrying no unnamed geometry
    must have grown no key at all (T-23-70's exact key-set equality)."""
    with _isolated_output_dir():
        report = generate_report("Baseline Proj", _baseline_extracted(), [])
        assert list(report.keys()) == PRE_23_TOP_LEVEL_KEYS
        assert "unnamed_geometry" not in report
        assert "run_settings" not in report

        with open(Path(report["_files"]["json"])) as f:
            assert "unnamed_geometry" not in json.load(f)


def test_an_empty_ledger_on_a_sheet_is_not_an_absent_one_it_is_simply_nothing():
    """A sheet that published a name for everything it measured writes an empty
    list. That must not conjure a key that says nothing."""
    extracted = _baseline_extracted()
    extracted[0]["unnamed_geometry"] = []
    with _isolated_output_dir():
        report = generate_report("Baseline Proj", extracted, [])
    assert list(report.keys()) == PRE_23_TOP_LEVEL_KEYS


def test_the_run_settings_key_records_the_ensemble_the_run_actually_used():
    """24-RESEARCH Open Question 1: neither saved paid run records the N that
    was in effect, and the environment is not an artifact. It need never be
    unanswerable again."""
    from bobbytailor.pipeline.takeoff_pipeline import RunSettings

    settings = RunSettings(
        deterministic_enabled=True,
        ensemble_n=3,
        learning_enabled=True,
        budget_usd=10.0,
        verify_retry_enabled=True,
        raster_force=False,
    )
    with _isolated_output_dir():
        report = generate_report(
            "Settings Proj", _baseline_extracted(), [],
            run_settings=settings.to_dict(),
        )
        assert report["run_settings"]["ensemble_n"] == 3
        assert report["run_settings"]["budget_usd"] == 10.0

        with open(Path(report["_files"]["json"])) as f:
            on_disk = json.load(f)
        assert on_disk["run_settings"]["ensemble_n"] == 3, (
            "the answer must be in the artifact, not only in memory"
        )


def test_the_run_settings_key_is_absent_when_no_caller_supplied_it():
    """A pre-24 caller is untouched — the same contract `coverage` documents."""
    with _isolated_output_dir():
        report = generate_report("Baseline Proj", _baseline_extracted(), [])
    assert "run_settings" not in report


def test_the_report_holds_its_own_copy_of_the_run_settings():
    settings = {"ensemble_n": 3, "budget_usd": 10.0}
    with _isolated_output_dir():
        report = generate_report(
            "Settings Proj", _baseline_extracted(), [], run_settings=settings,
        )
    settings["ensemble_n"] = 99
    assert report["run_settings"]["ensemble_n"] == 3


def test_the_unnamed_geometry_section_is_absent_when_there_is_nothing_to_say():
    """The empty contract, copied deliberately from `_coverage_lines`: a report
    without unnamed geometry renders byte-identically. The section does not
    exist rather than existing and saying zero."""
    from bobbytailor.pipeline.reporter import _unnamed_geometry_lines

    assert _unnamed_geometry_lines({}) == []
    assert _unnamed_geometry_lines({"unnamed_geometry": []}) == []
    assert _unnamed_geometry_lines({"unnamed_geometry": None}) == []


def test_the_summary_renders_the_unnamed_geometry_section_when_there_is():
    with _isolated_output_dir():
        report = generate_report(
            "Unnamed Proj", _two_sheets_with_unnamed_geometry(), [],
        )
        text = Path(report["_files"]["summary_txt"]).read_text()

    assert "UNIDENTIFIED GEOMETRY" in text
    # The count, stated as a number a reader can act on.
    assert "Candidates measured but not identified:" in text
    assert "SYM-2" in text and "WALL-TC4" in text and "ROOM-3" in text
    # In WORDS: what is missing is the identity, not the quantity and not the
    # object. INV-08's reason this is a section and not only a JSON key.
    assert "MEASURED" in text
    assert "absence of a NAME, never of an object" in text
    # And the section sits where the plan puts it — after COVERAGE, before the
    # scale block, so the two adjacent questions read as adjacent.
    assert text.index("UNIDENTIFIED GEOMETRY") < text.index("SCALE CONFIDENCE")


def test_the_unnamed_section_breaks_down_by_kind_with_its_measured_quantity():
    """A reader must be able to tell 'we measured 56 of these and could not name
    them' from 'these do not exist'."""
    extracted = _two_sheets_with_unnamed_geometry()
    source = {
        r["candidate_id"]: r
        for sheet in extracted for r in sheet["unnamed_geometry"]
    }
    with _isolated_output_dir():
        report = generate_report("Unnamed Proj", extracted, [])
        text = Path(report["_files"]["summary_txt"]).read_text()

    section = text.split("UNIDENTIFIED GEOMETRY")[1].split("SCALE CONFIDENCE")[0]
    assert "symbol group" in section
    assert "wall thickness class" in section
    assert "room" in section
    # Every quantity is rendered from the record, so the assertion cannot drift
    # from what the producer wrote. A whole count keeps no decimal point: a
    # count of repeated marks is a number of things, and `56.0 EA` would read
    # like a measurement that happened to land on an integer.
    assert "{:g} EA".format(source["SYM-2"]["quantity"]) in section
    assert "{:g} LF".format(source["WALL-TC4"]["quantity"]) in section
    assert "{:g} SF".format(source["ROOM-3"]["quantity"]) in section
    # The sheet each came from, and the reason naming failed.
    assert "A-101" in section and "A-102" in section
    assert "wall class unlabeled" in section


def test_a_record_with_no_measured_quantity_is_never_rendered_as_a_zero():
    """INV-08 inside this section too: a room whose SF was never bound to a
    scale has no number, which is a different fact from a measured zero."""
    from bobbytailor.vision.fusion import UNNAMED_ROOM_UNLABELED

    extracted = _baseline_extracted()
    extracted[0]["unnamed_geometry"] = [
        _ledger_record("ROOM-9", None, "SF", UNNAMED_ROOM_UNLABELED, "room"),
    ]
    with _isolated_output_dir():
        report = generate_report("Unnamed Proj", extracted, [])
        text = Path(report["_files"]["summary_txt"]).read_text()

    section = text.split("UNIDENTIFIED GEOMETRY")[1].split("SCALE CONFIDENCE")[0]
    assert "(not measured)" in section
    assert "no measured quantity" in section
    assert "0.0 SF" not in section, "a recorded absence is never summed as zero"


def test_the_individual_listing_is_capped_but_never_silently_truncated():
    """T-24-25: the cap bounds the page, never the account. A list that stops
    without saying it stopped is the same lie as a missing row."""
    from bobbytailor.pipeline.reporter import _UNNAMED_GEOMETRY_LIST_CAP
    from bobbytailor.vision.fusion import UNNAMED_NO_NAME_SOURCE

    over = _UNNAMED_GEOMETRY_LIST_CAP + 7
    extracted = _baseline_extracted()
    extracted[0]["unnamed_geometry"] = [
        _ledger_record(
            "SYM-%d" % i, 2, "EA", UNNAMED_NO_NAME_SOURCE, "symbol_group",
        )
        for i in range(over)
    ]
    with _isolated_output_dir():
        report = generate_report("Unnamed Proj", extracted, [])
        text = Path(report["_files"]["summary_txt"]).read_text()

    section = text.split("UNIDENTIFIED GEOMETRY")[1].split("SCALE CONFIDENCE")[0]
    assert "and 7 more" in section
    assert "takeoff.json" in section
    # The TOTAL is the whole population, not the printed page of it.
    assert str(over) in section.split("\n")[4]
    # And every one of them really is in the artifact the line points at.
    assert len(report["unnamed_geometry"]) == over


def test_an_unrecognised_record_kind_is_still_counted_and_shown():
    """A kind this renderer has never seen must not vanish from the section —
    a silent skip here is the defect the section exists to abolish."""
    from bobbytailor.vision.fusion import UNNAMED_NO_NAME_SOURCE

    extracted = _baseline_extracted()
    extracted[0]["unnamed_geometry"] = [
        _ledger_record(
            "FIX-1", 3, "EA", UNNAMED_NO_NAME_SOURCE, "fixture_cluster",
        ),
    ]
    with _isolated_output_dir():
        report = generate_report("Unnamed Proj", extracted, [])
        text = Path(report["_files"]["summary_txt"]).read_text()

    section = text.split("UNIDENTIFIED GEOMETRY")[1].split("SCALE CONFIDENCE")[0]
    assert "fixture cluster" in section
    assert "FIX-1" in section


def test_a_report_with_no_unnamed_geometry_still_renders_byte_for_byte():
    """The 23-17 parity pin, re-run: this plan adds nothing to a summary whose
    run named everything it measured."""
    with _isolated_output_dir():
        report = generate_report("Baseline Proj", _baseline_extracted(), [])
        text = Path(report["_files"]["summary_txt"]).read_text()
    assert _mask_timestamp(text) == PRE_23_SUMMARY_TXT
    assert "UNIDENTIFIED GEOMETRY" not in text


def _unnamed_report_data() -> dict:
    """A takeoff.json carrying a two-record ledger, as the reporter writes it."""
    from bobbytailor.vision.fusion import (
        UNNAMED_NO_NAME_SOURCE,
        UNNAMED_WALL_CLASS_UNLABELED,
    )

    first = _ledger_record(
        "SYM-6", 36, "EA", UNNAMED_NO_NAME_SOURCE, "symbol_group",
        evidence_texts=["A-101", "12"],
    )
    first["sheet"] = "A-101"
    second = _ledger_record(
        "WALL-TC4", 812.5, "LF", UNNAMED_WALL_CLASS_UNLABELED, "wall_class",
    )
    second["sheet"] = "A-102"
    return {
        "project_name": "P",
        "takeoff_summary": [{"item": "Doors", "quantity": 4, "unit": "EA"}],
        "unnamed_geometry": [first, second],
    }


def test_the_excel_workbook_carries_an_unidentified_geometry_sheet(tmp_path):
    """The estimator's own artifact must show what was measured but not named —
    otherwise the exclusion 24-04 made reads as 'not found on the drawings'."""
    wb = _export(_unnamed_report_data(), tmp_path)

    assert "Unidentified Geometry" in wb.sheetnames
    ws = wb["Unidentified Geometry"]
    # Row 1 is the caption, row 2 the header, one row per ledger record after.
    assert ws.max_row - 2 == 2

    rows = {
        ws.cell(row=r, column=1).value: [
            ws.cell(row=r, column=c).value for c in range(1, 8)
        ]
        for r in range(3, ws.max_row + 1)
    }
    assert set(rows) == {"SYM-6", "WALL-TC4"}

    sym = rows["SYM-6"]
    assert sym[1] == "symbol group"
    assert sym[2] == "A-101", "the sheet the candidate was measured on"
    assert sym[3] == 36, "the count is preserved exactly as measured"
    assert sym[4] == "EA"
    assert "A-101" in str(sym[5]), "the text found on the drawing"
    assert sym[6] == "no name source"

    assert rows["WALL-TC4"][3] == 812.5
    assert rows["WALL-TC4"][6] == "wall class unlabeled"


def test_the_excel_unidentified_sheet_says_measurements_not_absences(tmp_path):
    """T-24-23: a silently-omitted candidate read as 'does not exist' is the
    defect. The workbook must make that misreading impossible."""
    wb = _export(_unnamed_report_data(), tmp_path)

    caption = str(wb["Unidentified Geometry"]["A1"].value)
    assert "REAL MEASUREMENTS" in caption
    assert "deliberately excluded" in caption
    assert "NOT missing" in caption

    # And reachable from the run facts, for a reader who never opens the tab.
    about = {r[0].value: r[1].value for r in wb["About"].iter_rows(min_row=1)}
    assert "2 candidate(s)" in about["UNIDENTIFIED GEOMETRY"]
    assert "never missing from it" in about["UNIDENTIFIED GEOMETRY"]


def test_a_run_that_named_everything_gains_no_empty_excel_worksheet(tmp_path):
    """The same rule the summary section follows: nothing to report, no sheet."""
    wb = _export({
        "project_name": "P",
        "takeoff_summary": [{"item": "Doors", "quantity": 4, "unit": "EA"}],
        "unnamed_geometry": [],
    }, tmp_path)
    assert wb.sheetnames == ["Take-Off", "About"]
    assert "UNIDENTIFIED GEOMETRY" not in {
        r[0].value for r in wb["About"].iter_rows(min_row=1)
    }


def test_a_ledger_quantity_never_measured_stays_blank_in_the_excel_book(tmp_path):
    """INV-08 in the spreadsheet: a room whose SF was never bound to a scale has
    no number, and a zero here would be a number nobody measured."""
    from bobbytailor.vision.fusion import UNNAMED_ROOM_UNLABELED

    record = _ledger_record("ROOM-9", None, "SF", UNNAMED_ROOM_UNLABELED, "room")
    record["sheet"] = "A-101"
    wb = _export({
        "project_name": "P", "takeoff_summary": [],
        "unnamed_geometry": [record],
    }, tmp_path)
    ws = wb["Unidentified Geometry"]
    assert ws.cell(row=3, column=1).value == "ROOM-9"
    assert ws.cell(row=3, column=4).value is None, "blank, not zero"


def test_excel_drawing_text_beginning_with_an_equals_is_not_a_formula(tmp_path):
    """T-24-24. Measured, not assumed: openpyxl infers `data_type='f'` for a
    string starting with `=` and emits <f>…</f> — a live formula in a workbook
    built from words a draughtsman typed on a drawing."""
    from bobbytailor.vision.fusion import UNNAMED_NO_NAME_SOURCE

    record = _ledger_record(
        "SYM-1", 2, "EA", UNNAMED_NO_NAME_SOURCE, "symbol_group",
        evidence_texts=["=1+1", "@SUM(A1)"],
    )
    record["sheet"] = "=cmd|'/c calc'!A0"
    wb = _export({
        "project_name": "P", "takeoff_summary": [],
        "unnamed_geometry": [record],
    }, tmp_path)
    ws = wb["Unidentified Geometry"]

    for column in (3, 6):
        cell = ws.cell(row=3, column=column)
        assert cell.data_type == "s", (
            "drawing-derived text must be written as text, never as a formula"
        )
        assert str(cell.value).startswith("="), "and verbatim, not mangled"


def test_the_excel_exporter_makes_no_paid_call(tmp_path):
    """This whole export path is offline — it reads a saved run and nothing
    else. A workbook build must never be able to charge the API."""
    source = (Path(__file__).resolve().parents[1]
              / "scripts" / "export_takeoff_excel.py").read_text()
    for forbidden in ("ANTHROPIC_API_KEY", "analyze_drawing", "pytest.mark.golden"):
        assert forbidden not in source


def test_both_entry_points_hand_the_reporter_their_run_settings():
    """The 23-17 lesson in one grep: a reporter that CAN record the settings
    and a pipeline that never hands them over leaves the question exactly as
    unanswerable as before."""
    root = Path(__file__).resolve().parents[1]
    for rel in ("src/bobbytailor/pipeline/pdf_analyzer.py",
                "src/bobbytailor/scrape/scraper.py"):
        source = (root / rel).read_text()
        assert "run_settings=" in source, f"{rel} never passes run_settings"
